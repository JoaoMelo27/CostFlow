"""
Funções compartilhadas para localizar dados dentro das planilhas (.xlsb ou .xlsx).

Princípio central: NUNCA usar índice fixo de linha/coluna para achar "o mês
de agosto" ou "a coluna de RF". As planilhas-fonte não têm o mesmo layout
entre si (a mesma aba "Cash Cost 2026 Completo" tem os blocos de RF/Real
deslocados em 1 coluna entre os dois arquivos de custo). Por isso tudo aqui
localiza posições por:
  - texto do rótulo (nome da linha/coluna), ou
  - datas gravadas nas células de cabeçalho (serial do Excel).

É isso que faz o sistema funcionar em setembro, outubro etc. sem mexer no
código (regra do projeto: nada de valores fixos de um mês).
"""
from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from typing import Dict, List, Optional, Tuple

from pyxlsb import open_workbook

EXCEL_EPOCH = datetime(1899, 12, 30)

# Meses em português -> número. Usado para ler cabeçalhos de tabela dinâmica
# como "Soma de jan/26". Aceita sufixos extras (ex.: "jan/262"), um artefato
# que aparece em tabelas dinâmicas copiadas.
PT_MONTHS = {
    "jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6,
    "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12,
}


def excel_serial_to_date(serial: float) -> date:
    return (EXCEL_EPOCH + timedelta(days=serial)).date()


def is_plausible_date_serial(value) -> bool:
    """True se o número parece uma data do Excel entre ~2015 e ~2035."""
    return isinstance(value, (int, float)) and 42000 <= value <= 50000


@dataclass
class SheetGrid:
    """Uma aba inteira em memória, como dicionário esparso {linha: {coluna: valor}}.

    Simples e suficiente para estas planilhas (poucas centenas de linhas
    relevantes; abas maiores são lidas com max_rows para não carregar o que
    não interessa). Linhas e colunas começam em 0.
    """
    rows: Dict[int, Dict[int, object]]
    n_rows: int

    def cell(self, r: int, c: int):
        return self.rows.get(r, {}).get(c)

    def row(self, r: int) -> Dict[int, object]:
        return self.rows.get(r, {})

    def find_label_row(self, col: int, text: str, exact: bool = True) -> Optional[int]:
        """Primeira linha em que a célula (linha, col) tem o texto (igual, ou 'contém' se exact=False)."""
        for r, cols in self.rows.items():
            v = cols.get(col)
            if isinstance(v, str):
                if exact and v.strip() == text:
                    return r
                if not exact and text.lower() in v.lower():
                    return r
        return None

    def find_label_rows(self, col: int, text: str, exact: bool = True) -> List[int]:
        found = []
        for r, cols in self.rows.items():
            v = cols.get(col)
            if isinstance(v, str):
                if exact and v.strip() == text:
                    found.append(r)
                elif not exact and text.lower() in v.lower():
                    found.append(r)
        return sorted(found)


def _eh_xlsb(file_path: str) -> bool:
    """Decide o formato pelo CONTEÚDO, não pela extensão: .xlsb tem 'xl/workbook.bin'; .xlsx tem 'xl/workbook.xml'."""
    with zipfile.ZipFile(file_path) as z:
        nomes = z.namelist()
    if "xl/workbook.bin" in nomes:
        return True
    if "xl/workbook.xml" in nomes:
        return False
    raise ValueError("o arquivo não parece ser uma pasta de trabalho do Excel (.xlsx ou .xlsb).")


def _valor_xlsx(v):
    """Deixa o valor lido pelo openpyxl no mesmo formato do pyxlsb: datas viram serial do Excel."""
    if isinstance(v, bool):
        return None
    if isinstance(v, datetime):
        return (v - EXCEL_EPOCH).total_seconds() / 86400.0
    if isinstance(v, date):
        return float((datetime(v.year, v.month, v.day) - EXCEL_EPOCH).days)
    return v


def load_sheet(file_path: str, sheet_name: str, max_rows: Optional[int] = None) -> SheetGrid:
    """Lê uma aba de um .xlsb (pyxlsb) ou .xlsx (openpyxl) para o mesmo SheetGrid.

    Os índices de linha/coluna são 0-based nos dois formatos, então os leitores
    não precisam saber de qual formato o arquivo veio. No .xlsx são lidos os
    valores já calculados (data_only): o que o Excel gravou na última vez que salvou.
    """
    grid: Dict[int, Dict[int, object]] = {}
    n = 0
    if _eh_xlsb(file_path):
        with open_workbook(file_path) as wb:
            with wb.get_sheet(sheet_name) as sheet:
                for row in sheet.rows():
                    n += 1
                    if max_rows is not None and n > max_rows:
                        break
                    for c in row:
                        if c.v is not None:
                            grid.setdefault(c.r, {})[c.c] = c.v
        return SheetGrid(rows=grid, n_rows=n)

    from openpyxl import load_workbook
    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise KeyError(f"aba '{sheet_name}' (abas do arquivo: {', '.join(wb.sheetnames)})")
        ws = wb[sheet_name]
        ws.reset_dimensions()          # alguns arquivos gravam a dimensão errada; deixa o openpyxl varrer tudo
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            n += 1
            if max_rows is not None and n > max_rows:
                break
            for c, v in enumerate(row):
                v = _valor_xlsx(v)
                if v is not None and v != "":
                    grid.setdefault(r, {})[c] = v
    finally:
        wb.close()
    return SheetGrid(rows=grid, n_rows=n)


@dataclass
class ScenarioBlock:
    scenario: str                            # 'budget' | 'rf' | 'real'
    month_cols: Dict[Tuple[int, int], int]   # (ano, mes) -> coluna
    fy_col: int


def find_scenario_blocks_by_date_row(
    grid: SheetGrid,
    header_row_search_range: range,
    fy_label_row_offset_range: range,
    scenario_keywords: Dict[str, List[str]],
) -> Dict[str, ScenarioBlock]:
    """Localiza os blocos de colunas mensais (Budget/RF/Real) numa aba que grava a data
    de cada mês na própria linha de cabeçalho (como 'Cash Cost 2026 Completo').

    Passo a passo:
      1. A linha de cabeçalho é a que tem mais células com cara de data.
      2. As colunas com data são agrupadas em blocos contíguos (sem furo).
      3. O cenário de cada bloco é lido num texto perto da coluna seguinte ao
         bloco (a coluna do total anual) ou da primeira coluna do bloco,
         olhando algumas linhas acima. Isso cobre o rótulo exatamente em cima
         da coluna e o rótulo 1 coluna deslocado (célula mesclada).
    """
    best_row, best_count = None, -1
    for r in header_row_search_range:
        count = sum(1 for v in grid.row(r).values() if is_plausible_date_serial(v))
        if count > best_count:
            best_row, best_count = r, count
    if best_row is None or best_count == 0:
        raise ValueError("Não encontrei uma linha de cabeçalho com datas mensais.")

    header = grid.row(best_row)
    date_cols = sorted(c for c, v in header.items() if is_plausible_date_serial(v))

    blocks: List[List[int]] = []
    for c in date_cols:
        if blocks and c == blocks[-1][-1] + 1:
            blocks[-1].append(c)
        else:
            blocks.append([c])

    def label_near(col: int) -> Optional[str]:
        for dr in fy_label_row_offset_range:
            for probe_col in (col, col - 1, col + 1):
                v = grid.cell(best_row + dr, probe_col)
                if isinstance(v, str) and v.strip():
                    return v
        return None

    result: Dict[str, ScenarioBlock] = {}
    for block in blocks:
        fy_col = block[-1] + 1
        label_text = label_near(fy_col) or label_near(block[0]) or ""
        scenario = None
        for scen, keywords in scenario_keywords.items():
            if any(kw.lower() in label_text.lower() for kw in keywords):
                scenario = scen
                break
        if scenario is None:
            # bloco não identificado (ex.: um bloco de outro ano): fica de fora sem barulho
            continue
        month_cols = {}
        for c in block:
            d = excel_serial_to_date(header[c])
            month_cols[(d.year, d.month)] = c
        result[scenario] = ScenarioBlock(scenario=scenario, month_cols=month_cols, fy_col=fy_col)
    return result


def parse_pt_month_label(text: str) -> Optional[int]:
    """Número do mês em rótulos como 'Soma de jan/26' ou 'Soma de ago/262' (ignora o que vem depois da barra)."""
    m = re.search(r"\b([a-z]{3})/\d+", text.lower())
    if not m:
        return None
    return PT_MONTHS.get(m.group(1))


def months_up_to(ano: int, mes: int) -> List[Tuple[int, int]]:
    """[(ano, 1), (ano, 2), ..., (ano, mes)]: os meses que entram no YTD."""
    return [(ano, m) for m in range(1, mes + 1)]
