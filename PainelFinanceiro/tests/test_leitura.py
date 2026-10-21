"""
Leitura das planilhas de ponta a ponta (regressão).

Gera as planilhas fictícias de cada mês, roda o pipeline (o mesmo do botão
Atualizar) e compara com o gabarito calculado direto pelo modelo fictício.
Se algum leitor quebrar (coluna errada, YTD errado, cenário trocado), falha aqui.
"""
import os
import shutil
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import tests  # noqa: F401
import modelo_ficticio as mf
from planilhas_ficticias import gerar_planilhas, nomes_arquivos
from painel import pipeline
from painel.leitores.ocf import extract_ocf
from painel.leitores.planilha import excel_serial_to_date, is_plausible_date_serial, load_sheet


class TestLeituraDasPlanilhas(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.base = mf.gerar_base()
        cls.pastas = {}
        for mes in (1, 6, 8, 12):
            pasta = os.path.join(cls.tmp, f"2026-{mes:02d}")
            gerar_planilhas(pasta, mes, cls.base)
            cls.pastas[mes] = pasta

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_valores_batem_com_o_gabarito(self):
        for mes, pasta in self.pastas.items():
            arquivos = pipeline.localizar_arquivos(pasta)
            self.assertEqual(pipeline.detectar_periodo(arquivos), (2026, mes))
            resultado = pipeline.processar(arquivos, 2026, mes)
            self.assertFalse(resultado.tem_erro)
            esperado = mf.valores_esperados(mes, self.base)
            self.assertEqual(set(resultado.valores), set(esperado))
            for cod, por_visao in esperado.items():
                for visao, por_cen in por_visao.items():
                    for cen, valor in por_cen.items():
                        lido = resultado.valores[cod][visao][cen]
                        self.assertAlmostEqual(lido, valor, delta=1e-6 * max(1.0, abs(valor)),
                                               msg=f"{mes:02d} {cod}/{visao}/{cen}")

    def test_registros_tem_81_valores(self):
        resultado = pipeline.processar(pipeline.localizar_arquivos(self.pastas[8]), 2026, 8)
        self.assertEqual(len(resultado.registros()), 81)        # 9 indicadores automáticos x 3 visões x 3 cenários

    def test_ocf_recusa_mes_diferente_do_seletor(self):
        caminho = os.path.join(self.pastas[8], nomes_arquivos(8)["ocf"])
        with self.assertRaises(ValueError) as erro:
            extract_ocf(caminho, 2026, 7)
        self.assertIn("posicionado em 2026-08", str(erro.exception))

    def test_pasta_vazia_diz_o_que_falta(self):
        vazia = tempfile.mkdtemp()
        with self.assertRaises(FileNotFoundError) as erro:
            pipeline.localizar_arquivos(vazia)
        self.assertIn("não encontrado", str(erro.exception))

    def test_arquivo_ocf_protegido(self):
        pasta = tempfile.mkdtemp()
        for tipo, nome in nomes_arquivos(8).items():
            destino = os.path.join(pasta, nome)
            if tipo == "ocf":       # simula o arquivo com rótulo de sensibilidade (formato OLE2)
                Path(destino).write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 100)
            else:
                shutil.copy(os.path.join(self.pastas[8], nome), destino)
        with self.assertRaises(pipeline.ArquivoProtegidoError):
            pipeline.localizar_arquivos(pasta)

    def test_arquivo_corrompido_cita_o_nome(self):
        pasta = tempfile.mkdtemp()
        for nome in nomes_arquivos(8).values():
            shutil.copy(os.path.join(self.pastas[8], nome), os.path.join(pasta, nome))
        quebrado = os.path.join(pasta, nomes_arquivos(8)["custo_opex"])
        Path(quebrado).write_bytes(b"PK isto nao e um excel")
        arquivos = pipeline.localizar_arquivos(pasta)
        with self.assertRaises(pipeline.ArquivoComProblemaError) as erro:
            pipeline.processar(arquivos, 2026, 8)
        self.assertIn(nomes_arquivos(8)["custo_opex"], str(erro.exception))


class TestLeitorXlsx(unittest.TestCase):
    def test_mesmos_indices_e_datas_que_o_xlsb(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "BD"
        ws.cell(row=4, column=3, value=datetime(2026, 9, 1))
        ws.cell(row=5, column=4, value="(=) OCF")
        ws.cell(row=5, column=6, value=1234.5)
        caminho = os.path.join(tempfile.mkdtemp(), "teste.xlsx")
        wb.save(caminho)
        g = load_sheet(caminho, "BD")
        self.assertTrue(is_plausible_date_serial(g.cell(3, 2)))
        self.assertEqual(excel_serial_to_date(g.cell(3, 2)).isoformat(), "2026-09-01")
        self.assertEqual(g.cell(4, 3), "(=) OCF")
        self.assertEqual(g.cell(4, 5), 1234.5)
        self.assertIsNone(load_sheet(caminho, "BD", max_rows=4).cell(4, 3))
        with self.assertRaises(KeyError):
            load_sheet(caminho, "NaoExiste")


if __name__ == "__main__":
    unittest.main()
