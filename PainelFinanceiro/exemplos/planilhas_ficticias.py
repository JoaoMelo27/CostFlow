"""
Escreve as 3 planilhas de exemplo (.xlsx) de um fechamento, com o MESMO
layout das planilhas reais que o painel lê:

  1. "CUSTO TOTAL - Gestão de Custos OPEX ..."  -> aba "Cash Cost 2026 Completo"
  2. "Gestão de Custos OPEX ..."                -> aba "Cash Cost 2026 Completo"
                                                   + aba "Meta G&A sem nexa e sem corp"
  3. "OCF Report ..."                           -> aba "BD"

Os nomes de aba e os rótulos de linha precisam ser idênticos aos reais,
porque é por eles que os leitores (painel/leitores) acham os dados. Os
números são fictícios (ver modelo_ficticio.py).

Detalhes de layout reproduzidos de propósito, porque os leitores precisam
lidar com eles nas planilhas reais:
  - os dois arquivos de custo NÃO têm as colunas no mesmo lugar (no CUSTO
    TOTAL os blocos RF e Real ficam uma coluna mais à direita);
  - o mês de cada coluna é uma DATA na linha de cabeçalho;
  - no OCF, o mês de fechamento é uma célula-seletor ao lado do texto "MONTH".
"""
from __future__ import annotations

import os
from datetime import date
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import modelo_ficticio as mf
from modelo_ficticio import ANO, MESES, Cenario

AVISO = "DADOS FICTÍCIOS - gerados para fins acadêmicos. Não representam números de nenhuma empresa."
_NEGRITO = Font(bold=True)
_CABECALHO = PatternFill("solid", fgColor="FFE5D6")


def nomes_arquivos(mes: int) -> Dict[str, str]:
    """Nomes no mesmo padrão dos reais (o sistema acha cada arquivo pelo nome)."""
    ciclo = f"RF {mes}+{12 - mes}"
    return {
        "custo_total": f"CUSTO TOTAL - Gestão de Custos OPEX 26-27 - Unid Exemplo-{ciclo}.xlsx",
        "custo_opex": f"Gestão de Custos OPEX 26 - Unid Exemplo-{ciclo}.xlsx",
        "ocf": f"OCF Report Exemplo {ciclo}.xlsx",
    }


def _put(ws, linha: int, coluna: int, valor, negrito: bool = False, fundo: bool = False):
    """Escreve usando índices 0-based (os mesmos que os leitores usam)."""
    cel = ws.cell(row=linha + 1, column=coluna + 1, value=valor)
    if negrito:
        cel.font = _NEGRITO
    if fundo:
        cel.fill = _CABECALHO
    return cel


def _aba_leia_me(wb: Workbook) -> None:
    ws = wb.create_sheet("LEIA-ME")
    _put(ws, 0, 0, AVISO, negrito=True)
    _put(ws, 2, 0, "Arquivo gerado por exemplos/gerar_exemplos.py com o mesmo layout das planilhas reais.")


# --------------------------------------------------------------------------- aba "Cash Cost 2026 Completo"
def _aba_cash_cost(wb: Workbook, cenarios: Dict[str, Cenario], so_opex: bool, deslocamento: int) -> None:
    """deslocamento = colunas extras antes dos blocos RF/Real (0 no OPEX, 1 no CUSTO TOTAL)."""
    ws = wb.create_sheet("Cash Cost 2026 Completo")
    titulo = "Gestão de Custos OPEX" if so_opex else "Custo Total (OPEX + demais itens)"
    _put(ws, 0, 1, f"{titulo} - Unidade Exemplo - {ANO}", negrito=True)
    _put(ws, 1, 1, AVISO)

    linha_rotulo, linha_datas = 2, 3
    inicio = {"budget": 3, "rf": 17 + deslocamento, "real": 31 + deslocamento}
    rotulo = {"budget": f"BUDGET {ANO}", "rf": f"RF-1 {ANO}", "real": f"REAL {ANO}"}
    for nome, col0 in inicio.items():
        _put(ws, linha_rotulo, col0 + 12, rotulo[nome], negrito=True, fundo=True)   # rótulo em cima da coluna FY
        for m in MESES:
            _put(ws, linha_datas, col0 + m - 1, date(ANO, m, 1), fundo=True).number_format = "mmm/yy"
        _put(ws, linha_datas, col0 + 12, "FY", negrito=True, fundo=True)

    linhas: List[tuple] = []          # (rótulo, {cenario: (valores por mês, valor FY)}, negrito)

    def soma(serie):  # FY de valores = soma do ano
        return serie, sum(serie.values())

    def bloco(func):
        return {nome: func(c) for nome, c in cenarios.items()}

    linhas.append(("Produção Alimentada (t)", bloco(lambda c: soma(c.producao)), False))
    linhas.append(("Câmbio (R$/US$)", bloco(lambda c: (c.cambio, sum(c.cambio.values()) / 12)), False))
    linhas.append(("", {}, False))
    for nome_linha in mf.LINHAS_OPEX:
        linhas.append((nome_linha, bloco(lambda c, n=nome_linha: soma(c.linhas_opex[n])), False))
    linhas.append(("G&A (Despesas Gerais e Administrativas)", bloco(lambda c: soma(c.ga_total)), False))
    if not so_opex:
        for nome_linha in mf.LINHAS_FORA_OPEX:
            linhas.append((nome_linha, bloco(lambda c, n=nome_linha: soma(c.linhas_fora_opex[n])), False))
    custo = (lambda c: c.custo_opex) if so_opex else (lambda c: c.custo_total)
    linhas.append(("Custo Total (R$) | sem Efeito Corp com Nexa Energy", bloco(lambda c: soma(custo(c))), True))
    linhas.append(("", {}, False))

    def cash_cost(c):
        mensal = mf.cash_cost_usd_t(custo(c), c.cambio, c.producao)
        return mensal, mf.razao_ano(mf.custo_em_usd(custo(c), c.cambio), c.producao)

    linhas.append(("Custo Total (US$)", bloco(lambda c: soma(mf.custo_em_usd(custo(c), c.cambio))), False))
    linhas.append(("Cash Cost (US$/t Prod. Alimentada)", bloco(cash_cost), True))
    linhas.append(("Cash Cost (R$/t Prod. Alimentada)",
                   bloco(lambda c: ({m: custo(c)[m] / c.producao[m] for m in MESES},
                                    mf.razao_ano(custo(c), c.producao))), False))

    for i, (rot, valores, negrito) in enumerate(linhas):
        r = linha_datas + 2 + i
        if rot:
            _put(ws, r, 1, rot, negrito=negrito)
        for nome, (mensal, fy) in valores.items():
            col0 = inicio[nome]
            for m in MESES:
                _put(ws, r, col0 + m - 1, round(mensal[m], 6), negrito=negrito)
            _put(ws, r, col0 + 12, round(fy, 6), negrito=True)
    ws.column_dimensions["B"].width = 52


# --------------------------------------------------------------------------- aba "Meta G&A sem nexa e sem corp"
def _aba_ga(wb: Workbook, cenarios: Dict[str, Cenario]) -> None:
    """Três tabelas dinâmicas empilhadas: 'Budget 26', 'RF -1' e 'RF' (rótulo na coluna A)."""
    ws = wb.create_sheet("Meta G&A sem nexa e sem corp")
    _put(ws, 0, 0, AVISO, negrito=True)
    abrev = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]
    yy = ANO % 100
    linha = 2
    for rotulo, nome in ((f"Budget {yy}", "budget"), ("RF -1", "rf"), ("RF", "real")):
        c = cenarios[nome]
        _put(ws, linha, 0, rotulo, negrito=True)
        # o cabeçalho interno das tabelas reais diz "Soma de Real 2026" nos três blocos (vale o rótulo do bloco)
        _put(ws, linha + 1, 1, f"Soma de Real {ANO}")
        cab = linha + 2
        _put(ws, cab, 0, "Rótulos de Linha", negrito=True, fundo=True)
        for m in MESES:
            _put(ws, cab, m, f"Soma de {abrev[m - 1]}/{yy}", negrito=True, fundo=True)
        _put(ws, cab, 13, "Total Geral", negrito=True, fundo=True)
        r = cab + 1
        for centro, serie in c.linhas_ga.items():
            _put(ws, r, 0, centro)
            for m in MESES:
                _put(ws, r, m, round(serie[m], 2))
            _put(ws, r, 13, round(sum(serie.values()), 2))
            r += 1
        for rot, serie, total in (
            ("Total G&A (R$)", c.ga_total, sum(c.ga_total.values())),
            ("Produção Alimentada (t)", c.producao, sum(c.producao.values())),
            ("R$/Ton", {m: c.ga_total[m] / c.producao[m] for m in MESES}, mf.razao_ano(c.ga_total, c.producao)),
        ):
            _put(ws, r, 0, rot, negrito=True)
            for m in MESES:
                _put(ws, r, m, round(serie[m], 6), negrito=True)
            _put(ws, r, 13, round(total, 6), negrito=True)
            r += 1
        linha = r + 2
    ws.column_dimensions["A"].width = 30


# --------------------------------------------------------------------------- aba "BD" do OCF
def _aba_bd(wb: Workbook, cenarios: Dict[str, Cenario], mes: int) -> None:
    ws = wb.create_sheet("BD")
    _put(ws, 0, 0, "OCF Report - Unidade Exemplo", negrito=True)
    _put(ws, 1, 0, AVISO)
    H = 3                                     # linha com "MONTH" e os rótulos de cenário; datas na linha H + 1
    _put(ws, H, 2, "MONTH", negrito=True, fundo=True)
    _put(ws, H + 1, 2, date(ANO, mes, 1), fundo=True).number_format = "mmm/yy"   # célula-seletor do mês

    # linhas do relatório: (rótulo, função que devolve a série mensal de um cenário, tipo de agregação)
    def prod(c):
        return c.producao

    def nsr(c):
        return {m: c.receita_liquida[m] * 1000 / c.producao[m] for m in MESES}

    def margem_planilha(c):  # "OCF Margin" da planilha: sobre a receita BRUTA (o painel usa a líquida)
        return {m: c.ocf[m] / c.receita_bruta[m] for m in MESES}

    linhas = [
        ("Produção Alimentada (t)", prod, "soma"),
        ("Gross Revenues", lambda c: c.receita_bruta, "soma"),
        ("(-) Deductions", lambda c: {m: -(c.receita_bruta[m] - c.receita_liquida[m]) for m in MESES}, "soma"),
        ("(=) Net Revenues", lambda c: c.receita_liquida, "soma"),
        ("(-) Cash Cost", lambda c: {m: -c.custo_kusd[m] for m in MESES}, "soma"),
        ("(-) Other Operating Items", lambda c: {m: -c.outros_kusd[m] for m in MESES}, "soma"),
        ("(=) OCF", lambda c: c.ocf, "soma"),
        ("(=) OCF w/o Hedge", lambda c: c.ocf, "soma"),
        ("OCF Margin", margem_planilha, "razao_ocf_bruta"),
        ("NSR Financial", nsr, "razao_nsr"),
    ]
    primeira = 6
    for i, (rot, _, _) in enumerate(linhas):
        _put(ws, primeira + i, 3, rot, negrito=rot.startswith("(=)"))

    # ---- blocos mensais brutos: "ACT 2025" (só para lembrar que existem outros blocos), "BUD 2026", "RF ANT 2026", "RF 2026"
    blocos = [("ACT 2025", None), (f"BUD {ANO}", "budget"), (f"RF ANT {ANO}", "rf"), (f"RF {ANO}", "real")]
    col = 11
    for rotulo, nome in blocos:
        _put(ws, H, col, rotulo, negrito=True, fundo=True)
        ano_bloco = ANO - 1 if nome is None else ANO
        for m in MESES:
            _put(ws, H + 1, col + m - 1, date(ano_bloco, m, 1), fundo=True).number_format = "mmm/yy"
            for i, (_, func, _) in enumerate(linhas):
                if nome is None:
                    v = func(cenarios["budget"])[m] * 0.93          # ano anterior fictício
                else:
                    v = func(cenarios[nome])[m]
                _put(ws, primeira + i, col + m - 1, round(v, 6))
        col += 13

    # ---- tabela agregada MONTH / YTD / FY (no arquivo real são fórmulas sobre o seletor; aqui, valores)
    ate = [m for m in MESES if m <= mes]

    def agregado(func, tipo: str, c: Cenario, visao: str) -> float:
        meses = {"month": [mes], "ytd": ate, "fy": MESES}[visao]
        if tipo == "soma":
            serie = func(c)
            return sum(serie[m] for m in meses)
        if tipo == "razao_nsr":
            return sum(c.receita_liquida[m] for m in meses) * 1000 / sum(c.producao[m] for m in meses)
        return sum(c.ocf[m] for m in meses) / sum(c.receita_bruta[m] for m in meses)

    yy = f"{ANO % 100:02d}"
    col = 151
    for grupo, visao in (("MONTH", "month"), ("M+1", None), ("YTD", "ytd"), ("FY", "fy"), ("QUARTER", None)):
        _put(ws, H, col, grupo, negrito=True, fundo=True)
        for k, (prefixo, nome) in enumerate((("BUD", "budget"), ("RF ANT", "rf"), ("RF", "real"))):
            _put(ws, H + 1, col + k, f"{prefixo}´{yy}", negrito=True, fundo=True)
            if visao is None:
                continue
            for i, (_, func, tipo) in enumerate(linhas):
                _put(ws, primeira + i, col + k, round(agregado(func, tipo, cenarios[nome], visao), 6))
        col += 5
    ws.column_dimensions["D"].width = 30


# --------------------------------------------------------------------------- API
def gerar_planilhas(pasta: str, mes: int, base: mf.Base | None = None) -> Dict[str, str]:
    """Cria as 3 planilhas do fechamento de `mes` em `pasta`. Retorna {tipo: caminho}."""
    base = base or mf.gerar_base()
    cenarios = mf.cenarios_do_fechamento(base, mes)
    os.makedirs(pasta, exist_ok=True)
    nomes = nomes_arquivos(mes)
    caminhos = {k: os.path.join(pasta, v) for k, v in nomes.items()}

    for tipo in ("custo_total", "custo_opex"):
        wb = Workbook()
        wb.remove(wb.active)
        so_opex = tipo == "custo_opex"
        _aba_cash_cost(wb, cenarios, so_opex=so_opex, deslocamento=0 if so_opex else 1)
        if so_opex:
            _aba_ga(wb, cenarios)
        _aba_leia_me(wb)
        wb.save(caminhos[tipo])

    wb = Workbook()
    wb.remove(wb.active)
    _aba_bd(wb, cenarios, mes)
    _aba_leia_me(wb)
    wb.save(caminhos["ocf"])
    return caminhos
